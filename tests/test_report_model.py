import copy
import json
from pathlib import Path
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from pypdf import PdfReader
from src.data import report_model

ROOT=Path(__file__).resolve().parents[1]
ATLAS=ROOT/'output/flow_atlas.xlsx'
REFERENCE=ROOT/'inputs/1. Fichier et RAG/Flow_Atlas_presentation.pdf'
REPORT=ROOT/'output/report_01_32'


class ReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        wb=load_workbook(ATLAS,data_only=True,read_only=True)
        cls.tables={name:report_model.read_table(wb,name) for name in ('Atlas','Preuves','Questions','Interpretations')}
        wb.close()

    def build(self,change=None,dataset=None):
        tables=copy.deepcopy(self.tables)
        if change:change(tables)
        with patch.object(report_model,'read_table',side_effect=lambda wb,name:tables[name]):
            return report_model.build_model(ATLAS,dataset,REFERENCE)

    def metrics(self,model):return {m['metric_id']:m['metric_value'] for m in model['metrics']}

    def test_cost_and_chart_reconcile(self):
        model=self.build();m=self.metrics(model)
        self.assertEqual(m['active_cost'],1274)
        self.assertEqual(sum(v for k,v in m.items() if k.startswith('domain_')),m['active_cost'])
        self.assertEqual(sum(v for k,v in m.items() if k.startswith('decision_')),m['usage_count'])

    def test_input_change_updates_all_page_costs(self):
        def change(t):
            next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-01')['annual_cost_ke']+=10
        model=self.build(change);self.assertEqual(self.metrics(model)['active_cost'],1284)
        elements={e['element_id']:e for rows in model['pages'].values() for e in rows}
        self.assertEqual(elements['P5_KPI_07']['value'],1284)
        self.assertEqual(elements['P7_KPI_01']['value'],1284)
        self.assertIn('1 284',elements['P6_MONEY']['display_text'])

    def test_missing_cost_does_not_become_zero(self):
        def change(t):next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-01')['annual_cost_ke']=None
        model=self.build(change);m=self.metrics(model)
        self.assertIsNone(m['active_cost']);self.assertIsNone(m['stop_cost'])
        self.assertIsNone(m['domain_'+report_model.slug('Commerce')])

    def test_other_scope_and_company_are_supported(self):
        def change(t):
            for r in t['Atlas']:
                if r['record_type']=='usage':r['dataset']='other_company';r['organization']='Example Organisation'
        model=self.build(change,dataset='other_company')
        self.assertEqual(model['scope'],'Example Organisation')
        self.assertEqual(self.metrics(model)['active_count'],12)
        self.assertIn('Example Organisation',model['pages']['6'][0]['detail'])

    def test_duplicate_identity_rejected(self):
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            self.build(lambda t:t['Atlas'].append(copy.deepcopy(t['Atlas'][0])))

    def test_human_owner_lineage_and_raw_value(self):
        def change(t):
            row=next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-13')
            row['owner']='Reviewer-confirmed owner'
            row['human_validations']=json.dumps({'owner':{'value':row['owner']}})
            q=next(q for q in t['Questions'] if q['question_id']=='automobile:AI-13:owner')
            q.update(status='Human confirmed',answer=row['owner'],reviewer='Test reviewer',timestamp='2026-09-27T12:00:00Z',answer_evidence='Test validation')
        model=self.build(change)
        m=next(m for m in model['metrics'] if m['metric_id']=='missing_owners')
        self.assertEqual(m['metric_value'],4)
        self.assertIn('human:automobile:AI-13:owner',m['evidence_ids'])
        self.assertIsNone(next(r for r in model['usages'] if r['record_id']=='automobile:AI-13')['owner'])

    def test_missing_data_stays_missing(self):
        model=self.build();m=self.metrics(model)
        self.assertTrue(all(m[k] is None for k in ('coverage','skills','integration','alignment','savings')))
        self.assertTrue(all(r['profile_id']=='Unknown' for r in model['atlas'] if r['record_type']=='usage'))
        self.assertTrue(all(m['confidence_score'] is None for m in model['metrics']))

    def test_saved_report_and_traceability(self):
        report=json.loads((REPORT/'validation/report_validation.json').read_text(encoding='utf-8'))
        self.assertTrue(report['passed'])
        self.assertTrue(report['page6_regression_pixel_identical'])
        self.assertEqual(len(report['samples']),7)
        self.assertEqual(len(PdfReader(REPORT/'portaland_report.pdf').pages),32)
        wb=load_workbook(REPORT/'portaland_report_data.xlsx',data_only=True)
        self.assertEqual(len(wb.sheetnames),45)
        dates=[row[2] for row in wb['00_README'] if row[1].value=='Generated at (UTC)']
        self.assertEqual(len(dates),1)
        self.assertTrue(dates[0].is_date)
        chart_sheets={'03_REPORT_DASHBOARD','04_REFERENCE_CHARTS'}
        self.assertTrue(all(ws.freeze_panes=='A2' and len(ws.tables)==1 for ws in wb.worksheets if ws.title not in chart_sheets))
        self.assertEqual(sum(len(wb[s]._charts) for s in chart_sheets),6)
        self.assertEqual(wb['03_REPORT_DASHBOARD']['B9'].value,0)
        self.assertEqual(wb['03_REPORT_DASHBOARD']['B19'].value,0)
        self.assertIn(wb['04_REFERENCE_CHARTS']['B9'].value,(None,''))
        wb.close()

    def test_counterfactual_does_not_change_source_or_claims(self):
        model=self.build();m=self.metrics(model)
        self.assertEqual(m['active_cost'],1274)
        self.assertEqual(m['scenario_stop_cost'],951)
        self.assertAlmostEqual(m['scenario_stop_score'],57.2)
        self.assertEqual(m['scenario_stop_issues'],4)
        self.assertEqual(m['issue_count'],8)
        self.assertEqual({r['record_id']:r['status'] for r in model['atlas']},
                         {r['record_id']:r['status'] for r in self.tables['Atlas']})
        self.assertEqual(m['e7_candidates'],2)
        self.assertEqual(m['strategic_pilots'],2)
        self.assertEqual(m['valuable_ideas'],3)

    def test_scenario_missing_cost_propagates(self):
        def change(t):next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-01')['annual_cost_ke']=None
        m=self.metrics(self.build(change))
        self.assertIsNone(m['scenario_stop_cost'])
        self.assertIsNone(m['outside_cost'])
        self.assertIsNone(m['black_boxes'])

    def test_no_active_usages(self):
        def change(t):
            for r in t['Atlas']:
                if r['record_type']=='usage':r['status']='Idée'
            t['Interpretations']=[]
        model=self.build(change);m=self.metrics(model)
        self.assertIsNone(m['scenario_stop_score'])
        self.assertEqual(m['scenario_stop_cost'],0)
        self.assertEqual(m['outside_si'],0)
        self.assertEqual(next(r for r in model['pages']['11'] if r['kind']=='inventory')['label'],'Aucune IA active')

    def test_extended_traceability_and_domain_reconciliation(self):
        model=self.build();m=self.metrics(model)
        self.assertEqual(set(model['pages']),{str(p) for p in range(1,33)})
        self.assertEqual(sum(v for k,v in m.items() if k.startswith('outside_domain_')),m['outside_cost'])
        self.assertEqual(sum(v for k,v in m.items() if k.startswith('action_')),m['issue_count'])
        sources={r['evidence_id'] for r in model['evidence']}
        for page in range(1,33):
            for row in model['pages'][str(page)]:
                if row['data_status']!='MISSING':self.assertTrue(row['evidence_ids'])
                self.assertLessEqual(set(row['evidence_ids']),sources)

    def test_reference_examples_never_enter_business_metrics(self):
        model=self.build();m=self.metrics(model)
        self.assertEqual(m['active_cost'],1274)
        examples=model['illustrations']
        self.assertEqual(len(examples),33)
        self.assertEqual(next(r['value'] for r in examples if r['reference_page']==24 and r['category']=='Coût complet'),1282)
        self.assertTrue(all(r['data_role']=='REFERENCE_ILLUSTRATION' for r in examples))
        self.assertTrue(all(not any(e.startswith('illustration:') for e in r['evidence_ids']) for r in model['metrics']))
        actual=[r['value'] for r in examples if r['reference_page']==22 and r['series']=='Réel']
        self.assertEqual(actual,[38,44,52,None,None,None,None,None])
        self.assertIsNone(m['continuity']);self.assertIsNone(m['net_value']);self.assertIsNone(m['compliance'])

    def test_new_metrics_reconcile_and_unknown_decision_preserves_denominator(self):
        model=self.build();m=self.metrics(model)
        self.assertEqual(m['defensive']+m['maintain']+m['offensive']+m['other_decisions'],15)
        self.assertEqual(m['offensive_share'],20)
        self.assertEqual(sum(v for k,v in m.items() if k.startswith('stage_')),15)
        self.assertEqual(m['provider_top_cost'],433)
        self.assertAlmostEqual(m['provider_share'],433/1274*100)
        self.assertEqual(m['provider_price_shock'],866)
        self.assertAlmostEqual(m['classified_share'],11/12*100)
        def change(t):next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-01')['source_decision']='Unconfirmed'
        changed=self.metrics(self.build(change))
        self.assertEqual(changed['other_decisions'],1)
        self.assertEqual(changed['offensive_share'],20)

    def test_missing_provider_prevents_unfounded_concentration(self):
        def change(t):next(r for r in t['Atlas'] if r['record_id']=='automobile:AI-01')['provider']='Unknown'
        m=self.metrics(self.build(change))
        self.assertTrue(all(m[k] is None for k in ['provider_share','provider_top_cost','provider_top_count','provider_price_shock']))


if __name__=='__main__':unittest.main()

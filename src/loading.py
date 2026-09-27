import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def discover(folder):
    candidates = []
    for path in Path(folder).rglob("*.xlsx"):
        if path.name.startswith("~$"):
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        if {"Profils", "Usages IA"} <= set(wb.sheetnames):
            candidates.append(path)
        wb.close()
    if len(candidates) != 1:
        raise ValueError(f"Expected one Profils / Usages IA workbook, found {candidates}")
    return candidates[0]


def read_table(path, sheet, key, dataset):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows())]
    if key not in headers or len(set(headers)) != len(headers):
        raise ValueError(f"Invalid headers: {sheet}")
    records, evidence, seen = [], [], set()
    sha = digest(path)
    for row_num, cells in enumerate(ws.iter_rows(min_row=2), 2):
        if all(c.value is None for c in cells):
            continue
        raw = dict(zip(headers, [c.value for c in cells]))
        identity = raw[key]
        if not identity or identity in seen:
            raise ValueError(f"Missing or duplicate ID in {sheet}: {identity}")
        seen.add(identity)
        rid = f"{dataset}:{identity}"
        refs = {}
        for col_num, (header, cell) in enumerate(zip(headers, cells), 1):
            eid = f"{rid}:{header}"
            refs[header] = eid
            evidence.append(dict(evidence_id=eid, record_id=rid, field=header,
                                 value=cell.value, origin="source", file=str(path.resolve()),
                                 locator=f"{sheet}!{get_column_letter(col_num)}{row_num}", sha256=sha))
        records.append(dict(record_id=rid, dataset=dataset, raw=raw, evidence=refs))
    wb.close()
    return records, evidence


def rules(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    result = []
    sha = digest(path)
    for sheet in ("M6 Indices", "M7 Écarts", "M8 Décisions"):
        for row_num, cells in enumerate(wb[sheet].iter_rows(), 1):
            if any(c.value is not None for c in cells):
                result.append(dict(evidence_id=f"rule:{sheet}:{row_num}", record_id="rules",
                    field="rule", value=" | ".join(str(c.value) for c in cells if c.value is not None),
                    origin="reference", file=str(path.resolve()),
                    locator=f"{sheet}!A{row_num}:{get_column_letter(len(cells))}{row_num}", sha256=sha))
    wb.close()
    return result


def read_json(path, default):
    return json.loads(Path(path).read_text(encoding="utf-8")) if Path(path).exists() else default

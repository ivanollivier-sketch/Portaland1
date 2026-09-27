"""Independent reconciliation and limited structural comparison to PDF page 6."""
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from pypdf import PdfReader
from .loading import discover


def compare(inputs, output):
    source=load_workbook(discover(inputs),data_only=True,read_only=True)
    values=list(source['Usages IA'].values)
    records=[dict(zip(values[0],r)) for r in values[1:] if r[0]]
    active=[r for r in records if r['statut'] in ('Production','Pilote')]
    expected_cost=sum(r['cout_annuel_ke'] for r in active) if all(isinstance(r['cout_annuel_ke'],(int,float)) for r in active) else 'Unknown'
    source.close()
    wb=load_workbook(output/'flow_atlas.xlsx',data_only=True,read_only=True)
    observed=wb['Synthese']['B5'].value
    reference=next(Path(inputs).rglob('Flow_Atlas_presentation.pdf'))
    ref=PdfReader(reference).pages[5]
    generated=PdfReader(output/'flow_atlas_page6.pdf')
    page=generated.pages[0]
    text=page.extract_text()
    reftext=ref.extract_text()
    match=re.search(r'([\d\s]+)\s*k€?/an',reftext)
    ref_cost=int(re.sub(r'\s','',match.group(1))) if match else None
    checks={
        'one_page':len(generated.pages)==1,
        'same_page_dimensions':list(ref.mediabox)==list(page.mediabox),
        'independent_active_cost_matches':observed==expected_cost,
        'active_count_matches':wb['Synthese']['B4'].value==len(active),
        'required_sections':all(x in text for x in ('Capital cognitif','Urgence','Argent','Trois sujets')),
        'unknown_global_index':wb['Synthese']['B8'].value=='Unknown',
        'no_immediate_savings_claim':'réallouables dès maintenant' not in text,
    }
    wb.close()
    result=dict(checks=checks,passed=all(checks.values()),reference_active_cost_ke=ref_cost,
        generated_active_cost_ke=observed,
        difference_ke=observed-ref_cost if isinstance(observed,(int,float)) and ref_cost else None,
        comparison_scope='Dimensions, sections et rapprochement des chiffres. Revue visuelle manuelle du rendu PNG. Pas une équivalence des jeux de données.',
        expected_differences=['Indice de référence 131 non reproduit : données métier manquantes.',
            'Score de portefeuille déclaratif distinct du rendement financier.',
            'Recommandations issues du sous-ensemble actuel, proposées et non validées.',
            'Tri MVP : sévérité puis coût, un sujet par usage ; R8.3 non implémentée.'])
    (output/'comparison_page6.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    if not result['passed']:raise AssertionError(result)
    return result

"""Snapshot equality, sampled end-to-end lineage, and image diagnostics."""
import hashlib
import json
from pathlib import Path
import random
import shutil
import subprocess
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image,ImageChops,ImageStat
from ..data.report_model import read_table,STATUSES
from .data_book import PAGE_SHEETS


def validate_snapshot(model,output,reference,visual=False):
    output=Path(output);folder=output/'validation';folder.mkdir(exist_ok=True)
    wb=load_workbook(output/'portaland_report_data.xlsx',read_only=True,data_only=True)
    checks={}
    for p,sheet in PAGE_SHEETS.items():
        stored=read_table(wb,sheet)
        if len(stored)!=len(model['pages'][p]):raise AssertionError(f'Page row mismatch: {p}')
        for actual,expected in zip(stored,model['pages'][p]):
            for key,value in expected.items():
                got=actual.get(key)
                if isinstance(value,(dict,list)):got=json.loads(got)
                if value=='':value=None
                if got!=value:raise AssertionError(f"Page data mismatch {expected['element_id']}.{key}: {got!r} != {value!r}")
        checks[f'page_{p}_payload_equals_workbook']=True
    stored_metrics={r['metric_id']:r for r in read_table(wb,'09_REPORT_METRICS')}
    for m in model['metrics']:
        got=stored_metrics[m['metric_id']]['metric_value'];want=m['metric_value']
        if isinstance(want,(int,float)):
            assert isinstance(got,(int,float)) and abs(got-want)<1e-10
        else:assert got==want
    checks['all_metric_values_equal']=True
    stored_examples=read_table(wb,'44_REFERENCE_EXAMPLES')
    checks['reference_examples_match']=len(stored_examples)==len(model['illustrations']) and all(
        all(actual.get(k)==value for k,value in expected.items()) for actual,expected in zip(stored_examples,model['illustrations']))
    # Native charts read these formula cells; verify their cached results as well
    # as the canonical page tables, including a genuine zero and missing periods.
    dashboard=wb['03_REPORT_DASHBOARD']
    checks['dashboard_decisions_match']=all(dashboard[f'B{i}'].value==stored_metrics[mid]['metric_value'] for i,mid in enumerate(['defensive','maintain','offensive','other_decisions'],6))
    checks['dashboard_stages_match']=all(dashboard[f'B{i}'].value==m['metric_value'] for i,m in enumerate([m for m in model['metrics'] if m['metric_id'].startswith('stage_')],16))
    curve=wb['04_REFERENCE_CHARTS']
    series=list(dict.fromkeys(r['series'] for r in model['illustrations'] if r['reference_page']==22))
    checks['reference_chart_cells_match']=all(curve.cell(i+6,j+2).value in (None,'') if r['value'] is None else curve.cell(i+6,j+2).value==r['value']
        for j,s in enumerate(series) for i,r in enumerate([r for r in model['illustrations'] if r['reference_page']==22 and r['series']==s]))
    wb.close()
    pdf=PdfReader(output/'portaland_report.pdf')
    checks['page_count']=len(pdf.pages)==len(model['pages'])
    checks['page_dimensions']=all(tuple(p.mediabox)==(0,0,960,540) for p in pdf.pages)
    normalize=lambda s:' '.join(s.replace('œ','oe').split())
    checks['page_order']=all(normalize(model['pages'][str(i+1)][0]['display_text']) in normalize(p.extract_text()) for i,p in enumerate(pdf.pages))
    evidence={e['evidence_id']:e for e in model['evidence']};records={r['record_id'] for r in model['atlas']}
    elements=[r for rows in model['pages'].values() for r in rows]
    for r in elements:
        assert r['data_status'] in STATUSES
        assert set(r['evidence_ids'])<=set(evidence)
        assert set(r['source_record_ids'])<=records
    checks['all_element_links_resolve']=True
    illustrative_ids={r['evidence_id'] for r in model['illustrations']}
    checks['illustrations_excluded_from_business_metrics']=all(not(set(m['evidence_ids']) & illustrative_ids) for m in model['metrics'])
    checks['illustration_source_hashes_match']=all(evidence[x].get('file') and evidence[x].get('sha256')==hashlib.sha256(Path(evidence[x]['file']).read_bytes()).hexdigest() for x in illustrative_ids)
    rng=random.Random(731)
    candidates={
        'KPI':[r for r in elements if r['kind']=='kpi' and r['metric_ids'] and r['data_status']!='MISSING'],
        'chart':[r for r in elements if r['kind']=='chart'],
        'insight':[r for r in elements if r['kind']=='insight' and r['source_record_ids'] and r['evidence_ids'] and r['data_status']!='MISSING'],
    }
    samples=[];file_hashes={}
    for kind,count in [('KPI',3),('chart',2),('insight',2)]:
        for item in rng.sample(candidates[kind],count):
            page=item['element_id'].split('_')[0][1:];page_rows=model['pages'][page]
            sources=[]
            for eid in item['evidence_ids']:
                e=evidence[eid];file=e.get('file')
                if file and file not in file_hashes:
                    file_hashes[file]=hashlib.sha256(Path(file).read_bytes()).hexdigest() if Path(file).is_file() else None
                sources.append(dict(evidence_id=eid,source_file=file,locator=e.get('locator'),
                                    sha256_matches=bool(file and e.get('sha256') and file_hashes[file]==e['sha256'])))
            result=dict(kind=kind,pdf_element=item['element_id'],page_sheet=PAGE_SHEETS[page],page_row=page_rows.index(item)+2,
                        metric_ids=item['metric_ids'],atlas_record_ids=item['source_record_ids'],sources=sources,
                        status='PASS' if item['source_record_ids'] and sources and all(s['sha256_matches'] for s in sources) else 'INCOMPLETE')
            samples.append(result)
    checks['sampled_lineage_complete']=all(s['status']=='PASS' for s in samples)
    missing=[r['element_id'] for r in elements if r['data_status']=='MISSING']
    report=dict(checks=checks,passed=all(checks.values()),seed=731,samples=samples,
                missing_elements=missing,missing_note='Explicitly unavailable; no complete business calculation lineage can be asserted for missing referentials.',visual=[])
    if visual:
        poppler=shutil.which('pdftoppm')
        if not poppler:
            bundled=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/native/poppler/Library/bin/pdftoppm.exe'
            if bundled.exists():poppler=str(bundled)
        if not poppler:raise RuntimeError('pdftoppm is required for --validate')
        for p in map(int,model['pages']):
            reference_png=folder/f'reference_{p:02}.png';generated_png=folder/f'generated_{p:02}.png'
            subprocess.run([poppler,'-f',str(p),'-l',str(p),'-r','96','-png','-singlefile',str(reference),str(reference_png.with_suffix(''))],check=True,capture_output=True)
            subprocess.run([poppler,'-r','96','-png','-singlefile',str(output/'pages'/f'page_{p:02}.pdf'),str(generated_png.with_suffix(''))],check=True,capture_output=True)
            a=Image.open(reference_png).convert('RGB');b=Image.open(generated_png).convert('RGB')
            assert a.size==b.size
            difference=ImageChops.difference(a,b)
            similarity=1-sum(ImageStat.Stat(difference).mean)/(3*255)
            paired=Image.new('RGB',(a.width*2,a.height));paired.paste(a,(0,0));paired.paste(b,(a.width,0));paired.save(folder/f'comparison_{p:02}.png')
            report['visual'].append(dict(page=p,dimensions=a.size,pixel_similarity=similarity,
                 note='Diagnostic only; large shared background inflates similarity. Data correctness takes precedence.'))
        baseline=Path(__file__).resolve().parents[2]/'tmp/report_validation/page6_before.png'
        if baseline.exists():
            before=Image.open(baseline).convert('RGB');after=Image.open(folder/'generated_06.png').convert('RGB')
            report['page6_regression_pixel_identical']=before.size==after.size and ImageChops.difference(before,after).getbbox() is None
        previous=Path(__file__).resolve().parents[2]/'output/report_03_14/validation'
        if previous.exists():
            report['previous_pages_pixel_identical']={}
            for p in range(3,15):
                old=previous/f'generated_{p:02}.png'
                if old.exists():
                    before=Image.open(old).convert('RGB');after=Image.open(folder/f'generated_{p:02}.png').convert('RGB')
                    report['previous_pages_pixel_identical'][str(p)]=before.size==after.size and ImageChops.difference(before,after).getbbox() is None
    (folder/'report_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    lines=['# Report validation','',f"Result: {'PASS' if report['passed'] else 'FAIL'}",'',
           'Seven elements sampled with fixed random seed 731: three KPIs, two charts, two insights.',
           '', '| PDF element | Page worksheet / row | Metrics | Atlas IDs | Result |','|---|---|---|---|---|']
    for s in samples:
        lines.append(f"| {s['pdf_element']} | {s['page_sheet']} / {s['page_row']} | {', '.join(s['metric_ids']) or 'Rule interpretation / source fields'} | {', '.join(s['atlas_record_ids'])} | {s['status']} |")
    for s in samples:
        lines.extend(['',f"## {s['pdf_element']}: original sources",''])
        for source in s['sources']:
            lines.append(f"- `{source['evidence_id']}` → `{source['source_file']}` → `{source['locator']}`; hash verified: {source['sha256_matches']}.")
    lines.extend(['','## Missing elements','',', '.join(missing),'',report['missing_note']])
    if visual:
        lines.extend(['','## Visual comparison','','Reference and generated PNGs use 96 dpi. Paired PNGs show reference left, output right.'])
        for row in report['visual']:lines.append(f"- Page {row['page']}: mean pixel similarity {row['pixel_similarity']:.4f} (diagnostic only).")
        lines.append(f"- Existing page 6 pixel-identical: {report.get('page6_regression_pixel_identical','baseline unavailable')}.")
    (folder/'traceability_report.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    if not report['passed']:raise AssertionError(checks)
    return report
